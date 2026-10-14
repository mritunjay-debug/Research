import io
import re
import time
import html
import requests
import pandas as pd
import streamlit as st
from docx import Document
from rapidfuzz.fuzz import ratio

APP_TITLE = "Citation Verification Agent"
CROSSREF = "https://api.crossref.org/works"
HEADERS = {"User-Agent": "CitationVerificationAgent/1.0 (academic reference checking)"}

st.set_page_config(page_title=APP_TITLE, page_icon="📚", layout="wide")
st.title("📚 Citation Verification Agent")
st.write("Upload a Word document containing your reference list. The app checks reference metadata against Crossref and produces a review workbook and a Word reference report.")
st.info("A database miss is not proof that a paper is fake. Suggested matches are evidence for human review; uncertain references are not silently rewritten.")

def clean(s):
    s = html.unescape(s or "")
    s = re.sub(r"\\\[([^]]+)\\\]\((https?://[^)]+)\\\)", r"\1", s)
    s = re.sub(r"\\([.,&])", r"\1", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()

def split_refs(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    paras = [clean(p.text) for p in doc.paragraphs if clean(p.text)]
    refs = []
    for p in paras:
        if p.lower().strip(": ") in {"references", "bibliography", "reference list"}:
            continue
        # A single paragraph per reference is preferred; numbered references are retained.
        refs.append(p)
    if not refs:
        raise ValueError("No reference paragraphs were found. Check that the DOCX contains selectable text.")
    return refs

def parse_ref(ref):
    doi_match = re.search(r"(?:https?://(?:dx\.)?doi\.org/|\bdoi\s*:\s*)(10\.\d{4,9}/[^\s<>]+)", ref, re.I)
    doi = doi_match.group(1).rstrip(".,;)") if doi_match else ""
    year_match = re.search(r"\b(19|20)\d{2}[a-z]?\b", ref)
    year = year_match.group(0) if year_match else ""
    title = ""
    # Common author-year-title style: (2022). Title. Journal...
    if year_match:
        tail = ref[year_match.end():].lstrip(" .,:)")
        title = re.split(r"\.\s+(?=[A-Z])", tail, maxsplit=1)[0].strip(" .")
        # Prefer quoted title if present
        q = re.search(r'[“"]([^”"]{12,})[”"]', ref)
        if q:
            title = q.group(1)
    else:
        q = re.search(r'[“"]([^”"]{12,})[”"]', ref)
        if q:
            title = q.group(1)
    title = re.sub(r"\s+", " ", title)
    return {"DOI in reference": doi, "Year in reference": year, "Parsed title": title}

def crossref_get(params):
    try:
        r = requests.get(CROSSREF, params=params, headers=HEADERS, timeout=20)
        if r.status_code == 200:
            return r.json().get("message", {}).get("items", [])
    except requests.RequestException:
        pass
    return []

def norm_doi(d):
    return re.sub(r"^https?://(?:dx\.)?doi\.org/", "", d or "", flags=re.I).strip().rstrip(".,;")

def author_text(item):
    arr = item.get("author") or []
    names = []
    for a in arr:
        family = a.get("family", "")
        given = a.get("given", "")
        names.append((given + " " + family).strip())
    return "; ".join(names)

def item_year(item):
    for key in ("published-print", "published-online", "published", "issued", "created"):
        parts = item.get(key, {}).get("date-parts", [])
        if parts and parts[0]:
            return str(parts[0][0])
    return ""

def item_title(item):
    return (item.get("title") or [""])[0].strip()

def format_candidate(item):
    authors = item.get("author") or []
    author_parts = []
    for a in authors:
        family = a.get("family", "")
        given = a.get("given", "")
        initials = " ".join(x[0] + "." for x in re.findall(r"\b\w+", given))
        author_parts.append(f"{family}, {initials}".strip(", "))
    if len(author_parts) > 8:
        author_string = ", ".join(author_parts[:6]) + ", … " + author_parts[-1]
    else:
        author_string = ", ".join(author_parts)
    year = item_year(item)
    title = item_title(item)
    journal = (item.get("container-title") or [""])[0]
    volume = item.get("volume", "")
    issue = item.get("issue", "")
    page = item.get("page") or item.get("article-number", "")
    doi = norm_doi(item.get("DOI", ""))
    out = f"{author_string} ({year}). {title}. {journal}"
    if volume:
        out += f", {volume}"
    if issue:
        out += f"({issue})"
    if page:
        out += f", {page}"
    if doi:
        out += f". https://doi.org/{doi}"
    return out.strip()

def verify(ref):
    parsed = parse_ref(ref)
    doi = parsed["DOI in reference"]
    title = parsed["Parsed title"]
    year = parsed["Year in reference"]
    candidates = []
    method = ""
    if doi:
        candidates = crossref_get({"filter": "doi:" + doi, "rows": 3})
        method = "DOI lookup"
    if not candidates and title:
        candidates = crossref_get({"query.title": title, "rows": 5, "select": "DOI,title,author,published-print,published-online,published,issued,created,container-title,volume,issue,page,article-number,type"})
        method = "Title search"
    if not candidates and ref:
        # fallback uses the full citation, but does not equate no result with a fake reference
        candidates = crossref_get({"query.bibliographic": ref, "rows": 5})
        method = "Bibliographic search"
    best, best_score = None, 0
    for c in candidates:
        score = ratio((title or ref).lower(), item_title(c).lower()) / 100 if item_title(c) else 0
        if year and item_year(c) == year[:4]:
            score = min(1, score + 0.08)
        if doi and norm_doi(c.get("DOI", "")).lower() == doi.lower():
            score = 1.0
        if score > best_score:
            best, best_score = c, score
    if not best:
        return {
            "Original reference": ref, **parsed, "Decision": "MANUAL CORRECTION REQUIRED",
            "Match confidence": "", "Match method": method or "No database result",
            "Crossref title": "", "Crossref authors": "", "Crossref year": "",
            "Journal": "", "Crossref DOI": "", "Suggested reference": "",
            "Notes": "No Crossref match returned. This is not evidence that the paper is fabricated; verify through publisher, DOI resolver, or Google Scholar.",
            "Google Scholar search": "https://scholar.google.com/scholar?q=" + requests.utils.quote(title or ref)
        }
    cdoi = norm_doi(best.get("DOI", ""))
    cyear = item_year(best)
    ctitle = item_title(best)
    authors = author_text(best)
    journal = (best.get("container-title") or [""])[0]
    doi_match = not doi or doi.lower() == cdoi.lower()
    title_score = ratio((title or "").lower(), ctitle.lower()) / 100 if title else best_score
    year_match = not year or not cyear or year[:4] == cyear[:4]
    confident = (bool(doi) and doi_match and title_score >= .82) or (not doi and title_score >= .94 and year_match)
    decision = "TOTALLY CORRECTED" if confident else "MANUAL CORRECTION REQUIRED"
    notes = []
    if doi and not doi_match: notes.append("DOI differs from best Crossref result")
    if title and title_score < .94: notes.append("Title differs from Crossref result")
    if year and cyear and year[:4] != cyear[:4]: notes.append("Year differs; check online-first versus issue year")
    if not notes: notes.append("Crossref metadata matched the reference at the applied threshold" if confident else "Candidate found but evidence is not strong enough for automatic acceptance")
    return {
        "Original reference": ref, **parsed, "Decision": decision,
        "Match confidence": round(title_score, 3), "Match method": method,
        "Crossref title": ctitle, "Crossref authors": authors, "Crossref year": cyear,
        "Journal": journal, "Crossref DOI": cdoi,
        "Suggested reference": format_candidate(best),
        "Notes": "; ".join(notes),
        "Google Scholar search": "https://scholar.google.com/scholar?q=" + requests.utils.quote(title or ref)
    }

def make_docx(rows):
    doc = Document()
    doc.add_heading("Citation Verification Report", 0)
    doc.add_paragraph("References are checked against Crossref metadata. Manual-review items are retained as supplied and must be verified by the author.")
    for i, row in enumerate(rows, 1):
        doc.add_paragraph(f"{i}. " + (row["Suggested reference"] if row["Decision"] == "TOTALLY CORRECTED" and row["Suggested reference"] else row["Original reference"]))
        if row["Decision"] != "TOTALLY CORRECTED":
            p = doc.add_paragraph("[MANUAL CORRECTION REQUIRED] " + str(row["Notes"]))
            p.style = "Caption"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

uploaded = st.file_uploader("Upload reference list (.docx)", type=["docx"])
if uploaded:
    try:
        refs = split_refs(uploaded.getvalue())
        st.success(f"Extracted {len(refs)} reference paragraph(s).")
        with st.expander("Preview extracted references"):
            for i, ref in enumerate(refs, 1):
                st.write(f"{i}. {ref}")
        if st.button("Verify references", type="primary"):
            progress = st.progress(0)
            results = []
            status = st.empty()
            for i, ref in enumerate(refs):
                status.write(f"Checking reference {i+1} of {len(refs)}…")
                results.append(verify(ref))
                progress.progress((i+1)/len(refs))
                time.sleep(.08)
            df = pd.DataFrame(results)
            st.session_state["citation_results"] = df
            status.write("Verification complete.")
    except Exception as e:
        st.error(f"Could not read this DOCX: {e}")

if "citation_results" in st.session_state:
    df = st.session_state["citation_results"]
    a, b, c = st.columns(3)
    a.metric("References checked", len(df))
    b.metric("TOTALLY CORRECTED", int((df["Decision"] == "TOTALLY CORRECTED").sum()))
    c.metric("MANUAL CORRECTION REQUIRED", int((df["Decision"] == "MANUAL CORRECTION REQUIRED").sum()))
    st.dataframe(df[["Original reference", "Decision", "Crossref title", "Crossref year", "Crossref DOI", "Notes"]], use_container_width=True, hide_index=True)
    excel = io.BytesIO()
    with pd.ExcelWriter(excel, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Full Audit")
        df[df["Decision"] == "MANUAL CORRECTION REQUIRED"].to_excel(writer, index=False, sheet_name="Manual Review")
        df[df["Decision"] == "TOTALLY CORRECTED"].to_excel(writer, index=False, sheet_name="Matched")
    st.download_button("Download Excel audit (.xlsx)", excel.getvalue(), "citation_verification_audit.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.download_button("Download Word reference report (.docx)", make_docx(df.to_dict("records")), "corrected_reference_report.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    st.caption("This creates a separate reference-list report; it does not modify in-text citations in a full manuscript. Always review suggested metadata before submission.")
